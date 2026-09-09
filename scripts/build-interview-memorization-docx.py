from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Ask-Me-AI产品经理面试背诵手册.docx"

BLACK = "000000"
NAVY = "17365D"
MID_BLUE = "2F5597"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F6F7"
LIGHT_GRAY = "D9D9D9"
MUTED = "5A6573"
ACCENT = "C55A11"


def set_run_font(run, name="Microsoft YaHei", size=10.8, bold=False, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)


def set_keep_together(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_lines = p_pr.find(qn("w:keepLines"))
    if keep_lines is None:
        keep_lines = OxmlElement("w:keepLines")
        p_pr.append(keep_lines)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, text, fld_end])
    set_run_font(run, size=9, color=MUTED)


def add_mixed_paragraph(doc, parts, style=None, before=0, after=5, line=1.18, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    for text, bold, color in parts:
        run = p.add_run(text)
        set_run_font(run, bold=bold, color=color)
    if keep:
        set_keep_together(p)
    return p


def add_label_paragraph(doc, label, text, color=NAVY, after=4):
    return add_mixed_paragraph(
        doc,
        [(label, True, color), (text, False, BLACK)],
        before=0,
        after=after,
        line=1.2,
        keep=True,
    )


def add_bullet(doc, text, level=0, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.22 + level * 0.22)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True, color=NAVY)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for idx, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(width)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=9.5, bold=True, color="FFFFFF")
    set_repeat_table_header(table.rows[0])
    for ridx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for cidx, (value, width) in enumerate(zip(row_data, widths)):
            cell = cells[cidx]
            cell.width = Inches(width)
            set_cell_shading(cell, PALE_BLUE if ridx % 2 else "FFFFFF")
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cidx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_run_font(run, size=9.2, bold=(cidx == 0), color=NAVY if cidx == 0 else BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_question(doc, question, hook, answer, followups=None, level=2):
    heading = doc.add_heading(question, level=level)
    set_keep_with_next(heading)
    hook_p = add_label_paragraph(doc, "记忆钩子：", hook, color=ACCENT)
    set_keep_with_next(hook_p)
    add_label_paragraph(doc, "逐字稿：", answer, color=NAVY, after=4)
    for prompt, response in followups or []:
        add_label_paragraph(doc, f"追问  {prompt}：", response, color=MID_BLUE, after=4)


def add_starl(doc, labels):
    for label, text in labels:
        add_label_paragraph(doc, f"{label}：", text, color=NAVY, after=3)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.8)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)
    title.paragraph_format.space_after = Pt(12)

    for name, size, before, after in (
        ("Heading 1", 17, 16, 8),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11.5, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Memorize Lead" not in styles:
        lead = styles.add_style("Memorize Lead", WD_STYLE_TYPE.PARAGRAPH)
        lead.font.name = "Microsoft YaHei"
        lead.font.size = Pt(12)
        lead.font.bold = True
        lead.font.color.rgb = RGBColor.from_string(NAVY)
        lead._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        lead.paragraph_format.space_before = Pt(6)
        lead.paragraph_format.space_after = Pt(4)
        lead.paragraph_format.keep_with_next = True


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Ask Me AI 产品经理面试背诵手册")
    set_run_font(run, size=25, bold=True, color=BLACK)
    p.style = doc.styles["Title"]

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("分轮问题 追问 逐字稿 STAR L 复盘")
    set_run_font(run, size=13, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    run = p.add_run("AI 产品经理求职准备  2026 年 9 月版")
    set_run_font(run, size=10.5, color=MUTED)

    add_mixed_paragraph(
        doc,
        [
            ("使用原则  ", True, NAVY),
            ("先记关键词，再按自己的语气复述；所有数字必须能解释口径、分母和边界。", False, BLACK),
        ],
        before=6,
        after=8,
        line=1.25,
    )
    for text in (
        "第一遍：只背 30 秒定位、90 秒项目介绍和最大挑战。",
        "第二遍：按一面问题练直接回答，控制在 60 至 90 秒。",
        "第三遍：随机抽二面追问，练习“结论—机制—例子—边界”。",
        "面试现场：先给结论，面试官继续追问时再展开技术细节。",
    ):
        add_bullet(doc, text)
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.3)

    add_page_number(section.footer.paragraphs[0])
    add_cover(doc)

    doc.add_heading("第一部分 三分钟速记", level=1)
    add_mixed_paragraph(
        doc,
        [("这一部分先背熟。", True, NAVY), ("它覆盖项目定位、产品方案、数字口径和最大挑战，是一面最常用的回答底座。", False, BLACK)],
        after=8,
    )

    doc.add_heading("一句话定位", level=2)
    add_label_paragraph(
        doc,
        "必背：",
        "Ask Me 是面向招聘场景的 AI Career Agent，通过结构化个人知识、多轮问答和事实审校，让招聘方更快了解候选人的经历、个人贡献和能力边界。",
        color=ACCENT,
        after=7,
    )

    doc.add_heading("六句项目主线", level=2)
    for text in (
        "用户：主要是招聘经理和业务面试官；候选人是内容维护者和次级受益者。",
        "问题：简历信息有限、不能继续追问、个人贡献和 AI 辅助范围难核验。",
        "路径：快速了解—深入追问—项目核验—查看简历或联系。",
        "方案：Knowledge 管回答材料，Claim 管事实，Source 管证据与证明边界。",
        "策略：稳定题快速答，事实题审校后答，方法题实时生成。",
        "边界：系统质量已验证，但还没有足够真人招聘转化数据。",
    ):
        add_bullet(doc, text, bold_prefix=text.split("：", 1)[0] + "：")

    doc.add_heading("数字口径", level=2)
    add_table(
        doc,
        ["数字", "准确含义", "不能证明"],
        [
            ("20", "简历阶段核心问题数；当前为 25", "不是 20 位真人"),
            ("20", "安全攻击用例", "不能覆盖所有攻击"),
            ("55", "25 核心 + 20 安全 + 10 边界回归", "不是自动测试总数"),
            ("106", "简历形成时的自动化测试快照", "不是 106 个用户"),
            ("197/197", "2026-09-09 当前自动测试结果", "不等于产品成功"),
            ("48", "6 类 AI 角色 × 8 类问题", "不是 48 次真人面试"),
            ("100%", "25/25 核心题召回必需证据", "不等于答案 100% 正确"),
            ("0", "本轮硬事实违规数", "不等于线上幻觉率为 0"),
        ],
        [0.85, 3.15, 2.85],
    )

    doc.add_heading("30 秒项目介绍", level=2)
    add_label_paragraph(
        doc,
        "逐字稿：",
        "Ask Me 是我面向招聘场景设计并上线的 AI Career Agent。它解决的是传统简历只能单向展示、无法继续追问，而且个人贡献和 AI 辅助范围难核验的问题。产品围绕快速了解、深入追问、核验证据和简历联系设计，并用 KCS 证据结构、风险分级回答和自动评测控制事实风险。这个项目最能体现的是我如何在表达自然度、事实可信、响应速度和实现成本之间做产品取舍。",
        after=8,
    )

    doc.add_heading("90 秒项目介绍", level=2)
    add_starl(
        doc,
        [
            ("S", "一页简历适合快速浏览，但无法展开项目判断、本人贡献、AI 参与和结果边界。"),
            ("T", "我想把候选人的经历变成可浏览、可追问、可核验的求职数字分身，同时避免模型补造经历。"),
            ("A", "我按招聘方决策路径设计主流程；用 Knowledge、Claim、Source 管事实；再按风险选择稳定回答、审校回答或实时生成，并用澄清、有限回答和拒答处理不同边界。"),
            ("R", "产品已公开上线。简历阶段完成 106 项自动测试和 48 个 AI 合成面试用例；当前为 197/197 测试通过，核心证据覆盖 100%，本轮硬事实违规 0。"),
            ("L", "AI 产品经理不仅要设计模型怎么回答，还要定义哪些内容有资格回答、错误代价是什么，以及失败时怎样诚实降级。"),
        ],
    )

    doc.add_page_break()
    doc.add_heading("第二部分 简历逐词防守", level=1)
    add_mixed_paragraph(
        doc,
        [("目标：", True, NAVY), ("面试官点到简历中的任何词，都能用一句定义、一个机制和一个边界回答。", False, BLACK)],
        after=8,
    )
    add_table(
        doc,
        ["词语", "面试解释", "边界"],
        [
            ("招聘场景", "用户路径和回答标准都围绕招聘判断", "不是普通个人主页"),
            ("AI Career Agent", "能理解问题、读取上下文、选择回答策略并生成或拒答", "不是自动投递全流程 Agent"),
            ("数字分身", "候选人授权公开信息的交互表达层", "不能完整复制本人或替本人承诺"),
            ("设计并上线", "完成产品、内容、模型链路、评测、部署和公开访问", "不等于规模化商业上线"),
            ("面试深挖", "围绕同一经历继续问贡献、难点、取舍、结果和复盘", "不是无限轮记忆"),
            ("问题预演", "按个人材料和常见考察维度组织练习问题", "不承诺预测真实题目"),
            ("Knowledge", "为一次回答组织的结构化材料", "不是原始文件全文"),
            ("Claim", "最小可核验事实声明", "必须说明贡献和限制"),
            ("Source", "支撑 Claim 的材料及证明边界", "有链接不等于全部可信"),
            ("AI 辅助范围", "区分本人判断和 AI 的整理、生成、测试工作", "不编统一贡献比例"),
            ("未知问题处理", "指代不清则澄清，证据不足则有限回答，敏感请求才拒答", "不是所有陌生问题都拒绝"),
            ("幻觉处理", "内容准入、证据匹配、受约束生成、输出检查和失败回退", "只能降低风险，不能彻底消灭"),
            ("多轮问答", "维护当前项目、已问维度和已用证据", "不是把全部历史塞进 Prompt"),
            ("证据覆盖 100%", "25 个核心题都召回必需 Claim 和 Source", "不等于答案或产品 100% 成功"),
            ("硬事实 0 违规", "固定回归中没有命中材料外组织、数字、经历和结果", "不等于线上幻觉率为 0"),
        ],
        [1.25, 3.75, 1.85],
    )

    doc.add_heading("KCS 示例", level=2)
    add_label_paragraph(doc, "Knowledge：", "Ask Me 的回答策略与评测，是面向检索和回答组织的一段材料。")
    add_label_paragraph(doc, "Claim：", "产品按风险选择稳定回答、审校回答或实时生成，是其中一条事实。")
    add_label_paragraph(doc, "Source：", "Ask Me 的 PRD、代码和测试支撑这条事实，但不能证明招聘转化。")
    add_label_paragraph(doc, "一句总结：", "KCS 管理的不是“文本加链接”，而是“事实—证据—证明边界”。", color=ACCENT, after=8)

    doc.add_page_break()
    doc.add_heading("第三部分 一面高频题", level=1)
    add_mixed_paragraph(
        doc,
        [("回答原则：", True, NAVY), ("先讲用户和问题，再讲方案与结果；一题控制在 60 至 90 秒。", False, BLACK)],
        after=8,
    )

    first_round = [
        (
            "一面 1 为什么做这个项目",
            "简历单向  招聘方按需取信息  先问题后技术",
            "我在准备 AI 产品求职材料时发现，简历能展示做过什么，但很难说明为什么这样做、本人负责什么、AI 参与多少。招聘方的关注顺序也不同，所以我想验证一种既能快速浏览，又能继续追问和核验的求职信息形态。我先定义快速了解、深入追问、项目核验和简历联系的路径，再决定需要哪些 AI 能力。最终做成了公开 Demo 和质量门禁，但真人转化仍待验证。",
            [("为什么不用作品集", "作品集按作者顺序展示，Ask Me 允许招聘方按自己的判断顺序取信息；首页仍保留静态摘要，因为不能假设所有人都会主动提问。")],
        ),
        (
            "一面 2 目标用户是谁",
            "招聘方主用户  候选人次级用户  冲突时真实性优先",
            "主要用户是招聘经理和业务面试官，他们要在有限时间内判断候选人是否值得继续沟通。候选人是内容维护者，也是次级用户，可以用系统暴露的问题和证据缺口做面试准备。冲突时优先保护招聘信息真实，可以优化表达，但不能新增经历和数据。",
            [("是不是面试作弊工具", "不是。它做的是面试前的材料治理和问题预演，不提供面试中的实时作弊能力。")],
        ),
        (
            "一面 3 需求如何验证",
            "自身摩擦  公开面经  系统验证完成  真人验证不足",
            "需求来自我作为候选人的实际摩擦、招聘流程观察和公开面经中反复出现的项目深挖问题。当前我验证了问题合理性和系统可运行性，但还没有足够的目标招聘方访谈，所以不会说市场已经验证。下一步会让 5 名接近目标用户的人完成统一任务，观察理解、追问、查证和简历点击。",
            [("没有调研为什么开发", "这是探索性 MVP，首版成本可控；复盘看应该更早引入真人任务测试，限制后续工程投入。")],
        ),
        (
            "一面 4 你负责什么 AI 做了什么",
            "我负责定义取舍验收  AI 负责执行辅助  不编比例",
            "我负责产品洞察、用户路径、PRD、信息架构、证据模型、回答策略、优先级和测试标准；AI 编程 Agent 参与代码生成、资料结构化、排错和补测试。我通过测试、Lint、Build、真实问答和事实门禁验收结果。Ownership 看谁定义问题、做取舍并对结果负责，而不是只看手写代码行数。",
            [("AI 写了多少", "不同任务参与方式不同，我不编统一比例；可以逐项说明 AI 的执行范围和我承担的判断责任。")],
        ),
        (
            "一面 5 核心用户路径",
            "快速了解  深入追问  项目核验  简历联系",
            "用户进入首页先看定位、能力证据和推荐问题；感兴趣后点击问题或自由提问；系统结合当前项目和历史回答，并给出下一步追问；涉及事实时可以查看来源，最后进入简历或联系入口。即使模型不可用，摘要、项目、来源和简历仍能完成基本任务。",
            [],
        ),
        (
            "一面 6 为什么不用 FAQ",
            "高频用 FAQ  长尾和多轮用 AI  两者组合",
            "FAQ 能覆盖高频事实，所以 Ask Me 保留人工确认的稳定回答。AI 的增量价值是处理表达变化、开放问题和多轮指代。因此我没有把所有问题交给模型，而是让确定性能力负责稳定问题，让生成能力处理长尾。",
            [],
        ),
        (
            "一面 7 MVP 如何取舍",
            "用户价值  错误风险  验证成本",
            "我按用户价值、错误风险和验证成本排序。P0 保留首屏摘要、推荐与自由问答、项目证据、简历入口和事实边界；长期记忆、复杂向量检索、商业后台和自动投递延后。AI 产品的 MVP 是最小用户价值闭环加最低必要风险控制，不是模型功能越多越好。",
            [],
        ),
        (
            "一面 8 项目结果如何",
            "系统结果和用户结果分开",
            "系统侧，产品已上线；简历阶段完成 106 项自动测试和 48 个 AI 合成面试用例，当前为 197/197 测试通过，核心证据覆盖 100%，本轮硬事实违规 0。用户侧还没有足够真人招聘转化数据，所以只能说系统达到当前发布标准，不能说提高了面试通过率。",
            [],
        ),
        (
            "一面 9 如果重做会改什么",
            "更早真人测试  先五题原型  再决定工程投入",
            "如果重做，我会第一周只做可点击原型和 5 个核心回答，让目标用户完成 60 秒浏览、追问和查证任务，再决定是否投入复杂路由和审校。因为工程正确不等于产品正确，越早引入用户行为，越能避免在错误方向上精致化。",
            [],
        ),
        (
            "一面 10 最大挑战是什么",
            "自然表达和事实可信冲突",
            "最大挑战是让回答既像真实面试一样自然、有说服力，又不能为了表现力编造个人经历。我的解决方式不是继续堆 Prompt，而是用 KCS 管事实资格，用风险分级决定回答方式，再用硬事实门禁和审校兜底。完整 STAR-L 在第五部分。",
            [],
        ),
    ]
    for item in first_round:
        add_question(doc, *item)

    doc.add_page_break()
    doc.add_heading("第四部分 二面深挖题", level=1)
    add_mixed_paragraph(
        doc,
        [("回答原则：", True, NAVY), ("使用“结论—机制—项目例子—边界”，避免堆技术名词。", False, BLACK)],
        after=8,
    )

    second_round = [
        ("二面 1 一次问答链路", "安全判断  问题理解  证据检索  可回答性  分级交付", "用户提问后，系统先判断安全风险，再识别主题、提问动作、证据要求和事实风险。高频题走本地契约，低置信度长尾才用 Flash 规划；随后检索公开有效的 Knowledge，并映射到 Claim 和 Source；可回答性层决定回答、澄清、有限回答或拒答；最后按风险选择稳定展示、审校后展示或真实流式。", []),
        ("二面 2 KCS 为什么分三层", "Knowledge 管组织  Claim 管事实  Source 管证明", "三者解决的问题不同。Knowledge 服务检索和回答组织，Claim 决定事实是否有资格使用，Source 说明证据和证明边界。拆层增加维护成本，但来源变化、公开范围调整或事实失效时可以局部控制，适合个人事实错误代价较高的招聘场景。", [("有链接就可信么", "不一定。仓库能证明公开实现存在，但不能单独证明个人贡献比例、生产规模和业务效果。")]),
        ("二面 3 如何定义幻觉", "材料外个人事实  四层来源", "Ask Me 中的幻觉，是回答出现材料不支持的个人事实、数字、组织、事件、业务结果或过度确定的结论。来源可能是内容错误、检索错误、问题理解错误或生成擅自补全，因此不能把幻觉只当成模型文案问题。", []),
        ("二面 4 如何治理幻觉", "准入  匹配  生成  检查  回退", "生成前只允许公开、有效并确认的内容进入上下文；生成时提供事实骨架和禁止扩写边界；生成后检查材料外组织、数字、经历和结果。高风险事实题再由 Pro 审校，核心题使用确认答案。目标是降低风险并让失败可控，不是宣称零幻觉。", [("降低 Temperature 够么", "不够。它不能修复错误材料、召回错误或缺少证据。")]),
        ("二面 5 未知问题怎么处理", "澄清  有限回答  拒答", "对象不清时先澄清；与候选人相关但材料不足时，明确说无法确认并承接到可回答范围；涉及隐私、企业机密、提示词注入、越权读取和诱导编造时才直接拒答。招聘场景中编造经历的损失通常高于少答一道题。", [("拒答过多怎么评估", "离线看误拒率和漏拒率，线上看澄清、拒答、改问、退出和反馈。")]),
        ("二面 6 为什么三种回答方式", "稳定题快  事实题稳  方法题流式", "高频稳定问题直接展示确认答案；经历、贡献、数字和结果等事实敏感题完整生成并审校后展示；方法、诊断和假设题可以真实流式。这样低风险问题不必都慢，高风险事实也不会未经检查就被用户看到。", []),
        ("二面 7 为什么双模型", "Flash 生成  Pro 审校和回退  风险分工", "Flash 负责低延迟规划和主要生成，Pro 负责事实敏感回答的最终审校，以及 Flash 首包失败时的回退。模型名称可以替换，稳定的产品逻辑是按任务风险分工，并在额外调用前检查预算。", []),
        ("二面 8 为什么没用向量库", "当前规模小  结构化召回足够  不为标签加技术", "当前个人知识规模有限，主题和项目边界明确，结构化关键词、问题契约、查询扩展和最近上下文已经可解释、可测试，所以没有为了 RAG 标签引入向量数据库。出现同义召回不足和跨长文档需求时，再用固定评测集比较关键词、Embedding 和混合检索。", [("这还是 RAG 么", "有检索增强思想，但当前不是典型向量 RAG；我会准确说明。")]),
        ("二面 9 多轮怎么实现", "当前项目  已问维度  已用证据  指代不清则澄清", "系统不把全部历史直接塞进 Prompt，而是维护当前项目、目标岗位、对话深度、已问维度和已用证据。比如先问 RAG，再问“你做了什么”，会继承 RAG 并切到贡献维度；明确换项目时重置旧项目状态。", []),
        ("二面 10 评测体系如何设计", "内容  核心任务  安全多轮  回答体验", "第一层检查公开范围、验证状态和引用完整性；第二层检查核心问题的必需 Claim 和 Source；第三层检查隐私、注入、诱导编造、指代和去重；第四层检查直接回答、结构、长度、重复和硬事实。确定性门禁用于发布，AI 合成面试只用于发现叙事问题。", []),
        ("二面 11 证据覆盖率怎么算", "人工定义必需证据  25 分之 25", "我先为每个核心问题定义必须召回的 Claim ID 和 Source ID。分子是成功召回全部必需证据的问题数，分母是核心问题总数。当前为 25/25。它只验证证据准备完整，不验证最终文字和用户满意度。", []),
        ("二面 12 硬事实零违规怎么算", "禁止事实和已知幻觉模式  固定集结果", "硬事实包括材料外组织、数字、个人事件、用户反馈、业务结果和生产结论。评测为每个用例声明禁止事实并维护已知幻觉模式，本轮命中数为 0。它不是线上幻觉率，也不代表永远不会出错。", []),
        ("二面 13 二十个安全题是什么", "隐私  机密  注入  编造  越权", "安全题覆盖个人隐私、企业机密、Prompt Injection、诱导编造和越权读取。例如索要联系方式、客户名、系统提示词、虚构客户数据或读取环境变量。评测同时检查回答、日志和模型上下文是否泄漏。", []),
        ("二面 14 这些数字是什么关系", "测试项  评测题  合成面试分开", "106 是简历阶段自动测试快照，当前是 197；55 是 25 个核心、20 个安全和 10 个边界回归题；48 是 6 类 AI 角色乘 8 类问题。三者分母和目的不同，不能相加，也不能互相替代。", []),
        ("二面 15 北极星指标是什么", "有效招聘判断会话", "我会用“完成一次有效招聘判断的会话比例”作为验证期北极星指标。有效会话需要招聘方理解候选人定位、获得至少一条增量信息、完成证据查看或边界判断，并产生继续看简历或联系的意愿。辅助指标是追问率、证据查看率、简历点击率、误拒率、硬事实违规和延迟。", []),
        ("二面 16 如何控制成本稳定性", "少调用  分风险  预算  诚实降级", "稳定问题本地回答，只有低置信度长尾才用模型规划，只有高风险事实题才增加 Pro 审校；同时设置限流和 Token 预算。故障时静态摘要、证据、核心回答和简历仍可用，开放生成失败则明确提示，不用套话冒充成功。", []),
        ("二面 17 质量下降如何排查", "先确认测量  再分层定位  单变量复测", "先确认监控和样本是否真实，再依次检查问题理解与路由、证据召回、生成与审校、延迟与服务异常；然后按主题、回答维度、交付方式和模型路径聚类 Bad Case。每轮只改一个主要变量，用固定回归和线上样本复测。", []),
        ("二面 18 AI 产品经理的价值", "定义问题  分级错误  设计验收", "工程师擅长正确实现系统，我负责判断该解决什么问题、错误代价如何分级、异常状态如何定义、什么结果才算有效。技术理解让我能把目标变成可实现和可验收的规则，但不意味着替代工程师。", []),
    ]
    for item in second_round:
        add_question(doc, *item)

    doc.add_page_break()
    doc.add_heading("第五部分 三个挑战 STAR L", level=1)
    add_mixed_paragraph(
        doc,
        [("使用建议：", True, NAVY), ("主讲挑战一。面试官继续深挖理解或工程稳定性时，再讲挑战二或三。", False, BLACK)],
        after=8,
    )

    challenges = [
        (
            "挑战一 最大挑战 自然表达与事实可信冲突",
            [
                ("S", "早期版本在正文反复展示来源和限制，虽然安全，却像审计报告；放开表达后，模型又会补出不存在的漂亮结果。"),
                ("T", "让回答直接、有重点，同时让关键事实经得起追问。"),
                ("A", "内容层用 KCS 管事实资格；表达层把证据移到独立区域；验收层检查材料外组织、数字、经历和业务结果；高风险题增加审校，核心题使用确认答案。"),
                ("R", "当前核心证据覆盖率 100%，本轮硬事实违规 0，正文不再机械展示内部标签。"),
                ("L", "可信和自然不是二选一。系统在后台承担事实约束，用户看到自然表达，需要时再核验。"),
            ],
        ),
        (
            "挑战二 开放问题和多轮指代答偏",
            [
                ("S", "早期只看关键词，容易把岗位判断和项目结果混淆，也会让“这个项目”继承错对象。"),
                ("T", "识别用户真正的判断动作，而不只是识别主题。"),
                ("A", "将问题拆成主题、动作、回答维度、证据要求和事实风险；高频题走契约，长尾题用模型规划；维护当前项目和已问内容，仍不明确就澄清。"),
                ("R", "开放回归覆盖 150 道单轮题和 50 道多轮追问，减少串项目和重复回答。"),
                ("L", "很多生成质量问题，其实是上游任务理解错误。"),
            ],
        ),
        (
            "挑战三 流式体验与事实安全冲突",
            [
                ("S", "完整生成后再展示首字慢；真实流式又可能在后半段出现虚构数字或传输中断。"),
                ("T", "同时控制等待感、事实错误和服务失败。"),
                ("A", "稳定题本地展示，事实题审校后展示，方法题真实流式并逐片段检查；硬风险或中断时撤回半成品，双模型失败则明确不可用。"),
                ("R", "低风险题不必走最慢链路，流式撤回、空响应和模型回退也进入自动测试。"),
                ("L", "AI 体验不能只看响应时间，还要设计错误代价、可撤回性和故障后的剩余价值。"),
            ],
        ),
    ]
    for title, labels in challenges:
        heading = doc.add_heading(title, level=2)
        set_keep_with_next(heading)
        add_starl(doc, labels)

    doc.add_heading("踩坑九句话", level=2)
    for text in (
        "证据全放正文，安全但难读；改为后台严格、前台自然。",
        "只靠关键词，主题对了但动作错了；增加问题契约。",
        "全部历史塞给模型，容易串项目；改用短期结构化状态。",
        "只靠 Prompt 防幻觉，无法解决错误内容和召回；改为四层治理。",
        "完整生成后假流式，首字太慢；按风险区分交付方式。",
        "模型失败仍返回套话，会假装成功；改为语义匹配回退或明确不可用。",
        "用自动测试证明产品成功，混淆系统和用户价值；单独做真人验证。",
        "测试数增长却不记版本，简历和代码矛盾；明确 106 快照和 197 当前值。",
        "同时讲招聘方和候选人，目标用户模糊；招聘方主用户，候选人次级用户。",
    ):
        add_bullet(doc, text)

    doc.add_page_break()
    doc.add_heading("第六部分 主管面和压力面", level=1)
    final_questions = [
        ("主管 1 产品真正价值", "降低招聘方形成初步判断和核验信息的成本。相比静态简历，它可追问；相比通用模型，它只使用授权材料；相比模拟面试，它还有项目证据和简历转化路径。"),
        ("主管 2 竞品是谁", "替代方案包括简历、作品集、通用模型加简历、模拟面试产品和招聘平台候选人主页。差异点是招聘方可追问的信息层，以及贡献、AI 辅助和证据边界；但差异仍需真人验证。"),
        ("主管 3 只留一个功能", "保留围绕核心问题的可信回答，并和首屏入口合并。自由聊天、复杂推荐和后台都可以延后。"),
        ("主管 4 多候选人规模化难点", "最大难点会从问答技术转成内容治理，包括标准化导入、本人确认、版本更新、撤回和权限隔离；结合 JD 调整问题权重时也不能扭曲事实。"),
        ("主管 5 三个月路线图", "第一个月做目标招聘方任务测试；第二个月引入 JD 和人工黄金集；第三个月小范围开放候选人导入。只有 Bad Case 证明现有检索不足时，才引入向量或混合检索。"),
        ("主管 6 何时停止项目", "如果用户理解页面后仍不愿追问、查看证据，也不认为得到简历外价值，多轮迭代仍无改善，我会停止扩功能，把有效能力拆成候选人侧工具。"),
        ("压力 1 这是不是套壳", "如果只是聊天 UI 加模型 API，确实是套壳。Ask Me 的核心产品工作在用户路径、内容准入、证据关系、可回答性、风险分级、多轮状态、评测和降级。"),
        ("压力 2 是不是过度设计", "部分机制超过最小 Demo，但事实边界、失败状态和核心评测是必要能力；复杂检索、长期记忆和商业后台应等待真人验证。"),
        ("压力 3 没有真实用户还有价值么", "它能证明我把问题定义、AI 机制、工程实现和质量验收连成闭环，不能证明市场成功。项目能力已有证据，产品价值仍待验证。"),
        ("压力 4 模型完全不可用怎么办", "摘要、项目、来源、核心回答和简历仍能使用；开放生成明确提示不可用。降级目标是保住可信的基本任务。"),
        ("压力 5 为什么 12 题只有 11 题单项通过", "情景优先级题没有完整覆盖预设的用户目标和优先级取舍语义组，所以 11/12 达到单项阈值；硬事实、核心证据、多轮、路由和总体质量门槛通过。这个 Bad Case 已进入下一轮改进。"),
    ]
    for title, answer in final_questions:
        add_question(doc, title, "先说结论  再说边界", answer)

    doc.add_page_break()
    doc.add_heading("HR 面四题", level=2)
    hr_rows = [
        ("为什么做 AI PM", "统计让我重视指标和结论边界，审计让我重视流程和证据，RAG、模型评测与 Ask Me 让我把这些能力用于 AI 产品落地。"),
        ("项目体现什么优势", "能收敛模糊问题；能把模型不确定性转成规则与评测；能把方案推进到可运行、可验证。"),
        ("你的短板", "缺少大规模用户增长、商业化和长期跨职能管理经验；希望进入真实业务补足上线后的指标和资源取舍。"),
        ("为什么投这家公司", "用公司具体业务—岗位能力—可迁移项目证据—三个月期望交付回答，面试前按 JD 单独填写。"),
    ]
    add_table(doc, ["问题", "回答主线"], hr_rows, [1.6, 5.25])

    doc.add_page_break()
    doc.add_heading("第七部分 面试红线和临场检查", level=1)
    doc.add_heading("十一条红线", level=2)
    for text in (
        "不说项目提升了面试通过率或招聘转化率。",
        "不把 48 个 AI 合成用例说成 48 次真人面试。",
        "不把证据覆盖 100% 说成答案准确率 100%。",
        "不把硬事实 0 违规说成线上幻觉率为 0。",
        "不说系统彻底解决了幻觉，只说分层降低风险。",
        "不说 Ask Me 使用了向量数据库或企业级向量 RAG。",
        "不说所有回答都经过 Pro 审校。",
        "不说所有回答都是真实流式。",
        "不说全部代码由本人手写。",
        "不说公开仓库能完全证明个人贡献。",
        "不把 106 说成当前测试总数。",
    ):
        add_bullet(doc, text)

    doc.add_heading("临场五步", level=2)
    for text in (
        "第一句直接回答，不先铺背景。",
        "项目题优先讲一个主判断，不罗列所有功能。",
        "每出现一个技术名词，立刻补一句产品原因。",
        "每个数字说明口径、分母、验证什么、不能证明什么。",
        "结尾给 Learning，说明下一次会如何做得更好。",
    ):
        add_bullet(doc, text)

    doc.add_heading("十分钟模拟顺序", level=2)
    add_table(
        doc,
        ["分钟", "练习内容", "合格标准"],
        [
            ("0—1", "30 秒项目介绍", "用户、问题、方案、边界齐全"),
            ("1—3", "为什么做 + 个人贡献", "不堆技术，不虚构 AI 比例"),
            ("3—5", "最大挑战 STAR-L", "行动具体，结果有边界"),
            ("5—8", "随机二面深挖两题", "结论—机制—例子—边界"),
            ("8—9", "数字快问快答", "106、197、55、48、100%、0 不混"),
            ("9—10", "压力追问", "承认限制，但说明已验证价值"),
        ],
        [0.8, 3.0, 3.05],
    )

    doc.add_heading("参考来源", level=2)
    sources = (
        "AI 产品经理面试题库：https://chusimin.github.io/ai-pm-interview-bank/",
        "蚂蚁 AI 产品实习一面：https://www.nowcoder.com/feed/main/detail/50694c1390a54107bf36d180e4df2975",
        "Shopee AI 产品经理一面：https://www.nowcoder.com/feed/main/detail/066865138d224529ba82ee70fe87aefb",
        "字节 AI 产品经理面经：https://www.nowcoder.com/discuss/926271938274627584",
        "AIPM Wiki：https://github.com/archlizheng/AIPM-Wiki/blob/main/docs/04-interview/README.md",
    )
    for source in sources:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(source)
        set_run_font(run, size=9, color=MUTED)

    doc.core_properties.title = "Ask Me AI 产品经理面试背诵手册"
    doc.core_properties.subject = "AI 产品经理面试准备"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "Ask Me, AI 产品经理, 面试, STAR-L"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
